"""
Library-Upload -> Text-Extraktion -> Engine /ingest + Tenant-Bibliothek (WS3-DATA).

SCOPE v2 §2.5 + SCOPE v3 "Assets — Upload-Extraktion + Library".

Ablauf pro Upload:
  1. Text-Extraktion je Typ:
       pdf            -> pypdf
       docx           -> python-docx
       csv/txt/md/... -> stdlib (utf-8, errors=replace)
       xlsx           -> openpyxl
       .exe/unbekannt -> keine Extraktion (chars:0 + Hinweis), Datei wird aber angenommen
  2. Engine /ingest (env ENGINE_URL, default http://engine:8020) mit label+content+levels
     (default ["client"]) -> Inhalt fließt in Archiv/Memory.
  3. Datei-Metadaten persistent speichern (SQLite unter /data bzw. /seed-Volume),
     Roh-Datei optional unter <store>/files/{client_id}/.

Contract-Antwort /library/upload:
  {ok, file:{id,name,mime,size,chars,client_id,created_at}, ingested, chars, graph_delta}

Tenant-Trennung strikt: client_id filtert IMMER (list/delete/roh-Datei-Pfad).
Kein Crash bei kaputten/leeren Dateien (Stub-first, SCOPE §4).
"""
from __future__ import annotations

import io
import json
import mimetypes
import os
import re
import sqlite3
import tempfile
import time
import uuid
from pathlib import Path
from typing import Any, Optional

import httpx

_ENGINE_URL = os.getenv("ENGINE_URL", "http://engine:8020").rstrip("/")

# Persistenter Store: erster beschreibbarer Kandidat gewinnt (analog scrape.py).
# /data + /seed sind im docker-compose als Volume/Bind-Mount vorgesehen.
_STORE_CANDIDATES = [
    os.getenv("LIBRARY_DIR", ""),
    "/data/library",
    "/seed/library",
    "/app/seed/library",
    str(Path(__file__).resolve().parent / "seed" / "library"),
    "./seed/library",
]

# Plaintext-Formate, die direkt als utf-8 dekodiert werden.
_TEXT_EXT = {
    ".txt", ".md", ".markdown", ".csv", ".tsv", ".json",
    ".log", ".rst", ".yaml", ".yml", ".ini", ".xml", ".html", ".htm",
}

_MAX_INGEST_CHARS = int(os.getenv("LIBRARY_MAX_INGEST_CHARS", str(200_000)))


# --------------------------------------------------------------------------- store
def _resolve_store_dir() -> str:
    for cand in _STORE_CANDIDATES:
        if not cand:
            continue
        try:
            d = Path(cand)
            d.mkdir(parents=True, exist_ok=True)
            # Beschreibbarkeit testen
            probe = d / ".write_probe"
            with open(probe, "a", encoding="utf-8"):
                pass
            try:
                probe.unlink()
            except Exception:
                pass
            return str(d.resolve())
        except Exception:
            continue
    # Letzter Ausweg: temp (nicht persistent, aber crash-frei)
    d = Path(tempfile.gettempdir()) / "cnode-library"
    d.mkdir(parents=True, exist_ok=True)
    return str(d.resolve())


class LibraryStore:
    """SQLite-Metadaten-Store + optionale Roh-Datei-Ablage. Tenant-scoped."""

    def __init__(self) -> None:
        self.dir = _resolve_store_dir()
        self.db_path = str(Path(self.dir) / "library.db")
        self.files_dir = Path(self.dir) / "files"
        try:
            self.files_dir.mkdir(parents=True, exist_ok=True)
        except Exception:
            pass
        self._init_db()

    def _conn(self) -> sqlite3.Connection:
        conn = sqlite3.connect(self.db_path, timeout=10.0)
        conn.row_factory = sqlite3.Row
        return conn

    def _init_db(self) -> None:
        with self._conn() as conn:
            conn.execute(
                """
                CREATE TABLE IF NOT EXISTS files (
                    id          TEXT PRIMARY KEY,
                    name        TEXT NOT NULL,
                    mime        TEXT,
                    size        INTEGER,
                    chars       INTEGER,
                    client_id   TEXT NOT NULL,
                    created_at  TEXT,
                    kind        TEXT,
                    title       TEXT,
                    path        TEXT,
                    thread_id   TEXT
                )
                """
            )
            conn.execute(
                "CREATE INDEX IF NOT EXISTS idx_files_client ON files(client_id)"
            )
            # additive Migration für ältere DBs (sqlite kennt kein ADD COLUMN IF NOT EXISTS).
            try:
                conn.execute("ALTER TABLE files ADD COLUMN thread_id TEXT")
            except Exception:
                pass

    def add(self, meta: dict[str, Any]) -> None:
        with self._conn() as conn:
            conn.execute(
                """
                INSERT OR REPLACE INTO files
                    (id, name, mime, size, chars, client_id, created_at, kind, title, path, thread_id)
                VALUES (:id, :name, :mime, :size, :chars, :client_id, :created_at,
                        :kind, :title, :path, :thread_id)
                """,
                {
                    "id": meta["id"],
                    "name": meta["name"],
                    "mime": meta.get("mime"),
                    "size": meta.get("size", 0),
                    "chars": meta.get("chars", 0),
                    "client_id": meta["client_id"],
                    "created_at": meta.get("created_at"),
                    "kind": meta.get("kind"),
                    "title": meta.get("title"),
                    "path": meta.get("path"),
                    "thread_id": meta.get("thread_id"),
                },
            )

    def list(self, client_id: str, thread_id: Optional[str] = None) -> list[dict[str, Any]]:
        with self._conn() as conn:
            if thread_id:
                rows = conn.execute(
                    """
                    SELECT id, name, mime, size, chars, created_at, thread_id, title
                    FROM files WHERE client_id = ? AND thread_id = ?
                    ORDER BY created_at DESC, name ASC
                    """,
                    (client_id, thread_id),
                ).fetchall()
            else:
                rows = conn.execute(
                    """
                    SELECT id, name, mime, size, chars, created_at, thread_id, title
                    FROM files WHERE client_id = ?
                    ORDER BY created_at DESC, name ASC
                    """,
                    (client_id,),
                ).fetchall()
        return [dict(r) for r in rows]

    def get(self, file_id: str, client_id: str) -> Optional[dict[str, Any]]:
        with self._conn() as conn:
            row = conn.execute(
                "SELECT * FROM files WHERE id = ? AND client_id = ?",
                (file_id, client_id),
            ).fetchone()
        return dict(row) if row else None

    def delete(self, file_id: str, client_id: str) -> bool:
        """Metadaten + Roh-Datei entfernen. Nur eigener client_id (Tenant-Guard)."""
        meta = self.get(file_id, client_id)
        if not meta:
            return False
        with self._conn() as conn:
            conn.execute(
                "DELETE FROM files WHERE id = ? AND client_id = ?",
                (file_id, client_id),
            )
        raw_path = meta.get("path")
        if raw_path:
            try:
                Path(raw_path).unlink(missing_ok=True)
            except Exception:
                pass
        return True

    def count(self, client_id: Optional[str] = None) -> int:
        with self._conn() as conn:
            if client_id:
                row = conn.execute(
                    "SELECT COUNT(*) AS n FROM files WHERE client_id = ?",
                    (client_id,),
                ).fetchone()
            else:
                row = conn.execute("SELECT COUNT(*) AS n FROM files").fetchone()
        return int(row["n"]) if row else 0

    def clients(self) -> int:
        with self._conn() as conn:
            row = conn.execute(
                "SELECT COUNT(DISTINCT client_id) AS n FROM files"
            ).fetchone()
        return int(row["n"]) if row else 0


_STORE: Optional[LibraryStore] = None


def get_store() -> LibraryStore:
    global _STORE
    if _STORE is None:
        _STORE = LibraryStore()
    return _STORE


# --------------------------------------------------------------------------- extraction
def _guess_mime(filename: str, content_type: Optional[str]) -> str:
    if content_type and content_type not in ("application/octet-stream", ""):
        return content_type
    guessed, _ = mimetypes.guess_type(filename or "")
    return guessed or "application/octet-stream"


def _extract_pdf(raw: bytes) -> str:
    from pypdf import PdfReader  # requirements: pypdf

    reader = PdfReader(io.BytesIO(raw))
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            continue
    return "\n".join(pages).strip()


def _extract_docx(raw: bytes) -> str:
    from docx import Document  # requirements: python-docx

    doc = Document(io.BytesIO(raw))
    parts: list[str] = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    # Tabellen mitnehmen (häufig bei Angeboten/Verträgen)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()


def _extract_xlsx(raw: bytes) -> str:
    from openpyxl import load_workbook  # requirements: openpyxl

    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    lines: list[str] = []
    for ws in wb.worksheets:
        lines.append(f"# {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(c.strip() for c in cells):
                lines.append("\t".join(cells))
    try:
        wb.close()
    except Exception:
        pass
    return "\n".join(lines).strip()


def _extract_text(filename: str, raw: bytes, mime: str) -> tuple[str, str, str]:
    """Gibt (text, kind, note) zurück. kind = pdf|docx|xlsx|text|binary.

    Wirft nie — kaputte Dateien liefern leeren Text + Hinweis.
    """
    ext = os.path.splitext((filename or "").lower())[1]

    # PDF
    if ext == ".pdf" or "pdf" in mime:
        try:
            return (_extract_pdf(raw), "pdf", "")
        except Exception as e:  # noqa: BLE001
            return ("", "pdf", f"pdf_extract_failed:{type(e).__name__}")

    # DOCX (nur echtes OOXML .docx, nicht Legacy .doc)
    if ext == ".docx" or "wordprocessingml" in mime:
        try:
            return (_extract_docx(raw), "docx", "")
        except Exception as e:  # noqa: BLE001
            return ("", "docx", f"docx_extract_failed:{type(e).__name__}")

    # XLSX (OOXML)
    if ext == ".xlsx" or "spreadsheetml" in mime:
        try:
            return (_extract_xlsx(raw), "xlsx", "")
        except Exception as e:  # noqa: BLE001
            return ("", "xlsx", f"xlsx_extract_failed:{type(e).__name__}")

    # Plaintext (csv/txt/md/...) — auch endungslose Dateien best-effort dekodieren
    if ext in _TEXT_EXT or not ext or mime.startswith("text/"):
        try:
            return (raw.decode("utf-8", errors="replace").strip(), "text", "")
        except Exception as e:  # noqa: BLE001
            return ("", "text", f"text_decode_failed:{type(e).__name__}")

    # .exe / unbekanntes Binärformat -> annehmen, aber nicht extrahieren
    return ("", "binary", f"no_extractor_for:{ext or mime or 'unknown'}")


# --------------------------------------------------------------------------- helpers
def _norm_levels(levels: Any) -> list[str]:
    """Akzeptiert list, JSON-String, Komma-String. Default ['client']."""
    valid = ("client", "market", "mesh")
    if levels is None:
        return ["client"]
    if isinstance(levels, str):
        s = levels.strip()
        if not s:
            return ["client"]
        parsed: Any = None
        if s.startswith("["):
            try:
                parsed = json.loads(s)
            except Exception:
                parsed = None
        if parsed is None:
            parsed = [p for p in re.split(r"[,\s]+", s) if p]
        levels = parsed
    if not isinstance(levels, (list, tuple)) or not levels:
        return ["client"]
    out = [str(x).strip().lower() for x in levels if str(x).strip().lower() in valid]
    return out or ["client"]


def _summarize(text: str, limit: int = 800) -> str:
    return re.sub(r"\s+", " ", text or "").strip()[:limit]


def _title_from(filename: str, text: str) -> str:
    for line in (text or "").splitlines():
        line = line.strip().lstrip("#").strip()
        if len(line) >= 4:
            return line[:120]
    return os.path.basename(filename or "Dokument")


def _safe_ext(filename: str) -> str:
    ext = os.path.splitext((filename or ""))[1]
    return ext if re.fullmatch(r"\.[A-Za-z0-9]{1,8}", ext or "") else ""


def _store_raw(store: LibraryStore, file_id: str, client_id: str,
               filename: str, raw: bytes) -> Optional[str]:
    """Roh-Datei unter <store>/files/{client_id}/{id}{ext} ablegen (best-effort)."""
    try:
        # client_id defensiv säubern (kein Pfad-Traversal)
        safe_client = re.sub(r"[^A-Za-z0-9_.-]", "_", client_id or "default") or "default"
        target_dir = store.files_dir / safe_client
        target_dir.mkdir(parents=True, exist_ok=True)
        target = target_dir / f"{file_id}{_safe_ext(filename)}"
        target.write_bytes(raw)
        return str(target.resolve())
    except Exception:
        return None


# --------------------------------------------------------------------------- ingest
async def _engine_ingest(node: dict[str, Any]) -> dict[str, Any]:
    try:
        async with httpx.AsyncClient(timeout=15.0) as client:
            resp = await client.post(f"{_ENGINE_URL}/ingest", json=node)
            if resp.status_code == 200:
                data = resp.json()
                delta = data.get("graph_delta") if isinstance(data, dict) else None
                return {
                    "engine": "live",
                    "ingested": True,
                    "graph_delta": delta or {"nodes": [], "edges": []},
                }
    except Exception:
        pass
    # Fallback ohne Engine (Stub-first): lokales Delta bauen
    local_delta = {
        "nodes": [{
            "id": f"doc_{node['props'].get('file_id', uuid.uuid4().hex[:8])}",
            "label": node["label"],
            "type": "Document",
            "props": node["props"],
        }],
        "edges": [],
    }
    return {"engine": "fallback", "ingested": False, "graph_delta": local_delta}


# --------------------------------------------------------------------------- public API
async def ingest_file(
    filename: str,
    raw: bytes,
    client_id: str,
    content_type: Optional[str] = None,
    levels: Any = None,
    store_raw: bool = True,
    thread_id: Optional[str] = None,
) -> dict[str, Any]:
    """Upload verarbeiten: extrahieren -> Engine /ingest -> Metadaten persistieren."""
    client_id = (client_id or "default").strip() or "default"
    filename = filename or "upload.bin"
    mime = _guess_mime(filename, content_type)
    lvls = _norm_levels(levels)

    text, kind, note = _extract_text(filename, raw, mime)
    chars = len(text)
    title = _title_from(filename, text)
    file_id = uuid.uuid4().hex
    created_at = time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime())

    store = get_store()
    raw_path = _store_raw(store, file_id, client_id, filename, raw) if store_raw else None

    # Engine /ingest — Inhalt fließt in Archiv/Memory (nur wenn Text vorhanden)
    excerpt = _summarize(text)
    node = {
        "label": title,
        "type": "Document",
        "content": text[:_MAX_INGEST_CHARS],
        "props": {
            "file_id": file_id,
            "filename": filename,
            "mime": mime,
            "kind": kind,
            "bytes": len(raw),
            "chars": chars,
            "excerpt": excerpt,
            "client_id": client_id,
            "levels": lvls,
        },
        "links": [],
        "client_id": client_id,
        "levels": lvls,
    }

    if chars > 0:
        ing = await _engine_ingest(node)
    else:
        # Keine Extraktion (.exe/unbekannt/kaputt): kein Ingest, aber annehmen + Metadaten
        ing = {"engine": "skipped", "ingested": False,
               "graph_delta": {"nodes": [], "edges": []}}

    meta = {
        "id": file_id,
        "name": filename,
        "mime": mime,
        "size": len(raw),
        "chars": chars,
        "client_id": client_id,
        "created_at": created_at,
        "kind": kind,
        "title": title,
        "path": raw_path,
        "thread_id": thread_id,
    }
    try:
        store.add(meta)
    except Exception:
        pass

    return {
        "ok": True,
        "file": {
            "id": file_id,
            "name": filename,
            "mime": mime,
            "size": len(raw),
            "chars": chars,
            "client_id": client_id,
            "created_at": created_at,
        },
        "ingested": bool(ing.get("ingested")),
        "chars": chars,
        "graph_delta": ing.get("graph_delta", {"nodes": [], "edges": []}),
        "kind": kind,
        "title": title,
        "levels": lvls,
        "engine": ing.get("engine"),
        "note": note or None,
        # Auszug des extrahierten Textes → der Chat kann den Inhalt referenzieren/einordnen.
        "excerpt": excerpt,
    }


def list_files(client_id: str, thread_id: Optional[str] = None) -> dict[str, Any]:
    client_id = (client_id or "default").strip() or "default"
    return {"files": get_store().list(client_id, thread_id or None), "client_id": client_id}


def delete_file(file_id: str, client_id: str) -> dict[str, Any]:
    client_id = (client_id or "default").strip() or "default"
    deleted = get_store().delete(file_id, client_id)
    return {"ok": deleted, "deleted": deleted, "id": file_id, "client_id": client_id}


def status() -> dict[str, Any]:
    store = get_store()
    return {
        "files": store.count(),
        "clients": store.clients(),
        "store": store.dir,
        "db": store.db_path,
    }
